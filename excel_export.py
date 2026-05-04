"""
excel_export.py — Tien ich xuat file Excel cho phan mem Quan ly Diem DH
"""
import os
from datetime import datetime
from tkinter import filedialog, messagebox

try:
    import openpyxl
    from openpyxl.styles import (Font, PatternFill, Alignment,
                                  Border, Side, GradientFill)
    from openpyxl.utils import get_column_letter
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False


# ── Style helpers ──────────────────────────────────────────────────────────────

def _thin_border():
    s = Side(style='thin', color='CCCCCC')
    return Border(left=s, right=s, top=s, bottom=s)

def _header_fill():
    return PatternFill('solid', fgColor='4F46E5')   # indigo

def _alt_fill():
    return PatternFill('solid', fgColor='F1F5FB')

def _center(wrap=False):
    return Alignment(horizontal='center', vertical='center', wrap_text=wrap)

def _left():
    return Alignment(horizontal='left', vertical='center')

def _set_col_width(ws, col, width):
    ws.column_dimensions[get_column_letter(col)].width = width

def _title_row(ws, text, ncols, row=1):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name='Calibri', bold=True, size=14, color='FFFFFF')
    c.fill = _header_fill()
    c.alignment = _center()
    ws.row_dimensions[row].height = 32

def _sub_row(ws, text, ncols, row=2):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name='Calibri', italic=True, size=10, color='64748B')
    c.alignment = _center()
    ws.row_dimensions[row].height = 18

def _write_header_row(ws, headers, row, height=24):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = Font(name='Calibri', bold=True, size=10, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='3730A3')
        c.alignment = _center(wrap=True)
        c.border = _thin_border()
    ws.row_dimensions[row].height = height

def _write_data_row(ws, values, row, alt=False):
    fill = _alt_fill() if alt else PatternFill('solid', fgColor='FFFFFF')
    for col, v in enumerate(values, 1):
        c = ws.cell(row=row, column=col, value=v)
        c.font = Font(name='Calibri', size=10)
        c.fill = fill
        c.border = _thin_border()
        c.alignment = _left() if col == 2 else _center()

def _footer(ws, row, ncols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1,
                value=f"Xuat ngay: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}   |   Phan mem Quan ly Diem DH v5.0")
    c.font = Font(name='Calibri', italic=True, size=9, color='94A3B8')
    c.alignment = _center()
    ws.row_dimensions[row].height = 16

def _ask_save(default_name):
    path = filedialog.asksaveasfilename(
        defaultextension='.xlsx',
        filetypes=[('Excel Workbook', '*.xlsx'), ('All files', '*.*')],
        initialfile=default_name,
        title='Luu file Excel'
    )
    return path

def _open_file(path):
    try:
        os.startfile(path)
    except Exception:
        pass


# ── 1. Bang diem ca nhan (Sinh vien) ──────────────────────────────────────────

def export_bang_diem_ca_nhan(sv_info, rows, hk_filter='Tat ca'):
    """
    sv_info : (ma_sv, ho_ten, ten_lop, ngay_sinh, gioi_tinh)
    rows    : list of (ma_mh, ten_mh, so_tc, cc, gk, ck, tb, chu, he4, trang_thai)
    """
    if not OPENPYXL_OK:
        messagebox.showerror('Loi', 'Chua cai openpyxl. Chay: pip install openpyxl'); return

    path = _ask_save(f"BangDiem_{sv_info[0]}.xlsx")
    if not path: return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Bang diem'

    # Title
    ncols = 10
    _title_row(ws, 'BANG DIEM HOC TAP', ncols, row=1)
    _sub_row(ws, f'Hoc ky: {hk_filter}   |   Xuat ngay: {datetime.now().strftime("%d/%m/%Y")}', ncols, row=2)

    # Student info block
    info_labels = [
        ('Ma sinh vien:', sv_info[0]),
        ('Ho va ten:',    sv_info[1]),
        ('Lop:',          sv_info[2] or '---'),
        ('Ngay sinh:',    sv_info[3] or '---'),
    ]
    for i, (lbl, val) in enumerate(info_labels):
        r = 3 + i
        c1 = ws.cell(row=r, column=1, value=lbl)
        c1.font = Font(bold=True, size=10, color='4F46E5')
        c2 = ws.cell(row=r, column=2, value=val)
        c2.font = Font(size=10)
        ws.row_dimensions[r].height = 18

    # Blank separator
    ws.row_dimensions[7].height = 6

    # Header row
    headers = ['STT', 'Ten mon hoc', 'So TC', 'CC', 'Giua ky', 'Cuoi ky',
               'Trung binh', 'Xep loai', 'He 4', 'Trang thai']
    _write_header_row(ws, headers, row=8)

    # Data rows
    gpa_sum = 0.0; tc_sum = 0
    for idx, r in enumerate(rows):
        dr = 9 + idx
        ma_mh, ten_mh, tc, cc, gk, ck, tb, chu, he4, tt = r
        vals = [idx+1, ten_mh, tc,
                cc if cc is not None else '-',
                gk if gk is not None else '-',
                ck if ck is not None else '-',
                f'{tb:.2f}' if tb else '-',
                chu or '-',
                f'{he4:.1f}' if he4 else '-',
                tt or '-']
        _write_data_row(ws, vals, dr, alt=(idx % 2 == 1))
        if he4 is not None and tc:
            gpa_sum += he4 * tc
            tc_sum  += tc

    # Summary row
    sr = 9 + len(rows)
    ws.merge_cells(start_row=sr, start_column=1, end_row=sr, end_column=2)
    c = ws.cell(row=sr, column=1, value='TONG KET')
    c.font = Font(bold=True, size=10, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='4F46E5')
    c.alignment = _center()

    gpa = round(gpa_sum / tc_sum, 2) if tc_sum else 0
    summary_vals = ['', '', tc_sum, '', '', '', '', 'GPA:', gpa, '']
    for col, v in enumerate(summary_vals[2:], 3):
        sc = ws.cell(row=sr, column=col, value=v)
        sc.font = Font(bold=True, size=10)
        sc.fill = PatternFill('solid', fgColor='E0E7FF')
        sc.alignment = _center()
        sc.border = _thin_border()
    ws.row_dimensions[sr].height = 22

    # Column widths
    widths = [5, 32, 6, 6, 8, 8, 10, 8, 6, 12]
    for i, w in enumerate(widths, 1):
        _set_col_width(ws, i, w)

    _footer(ws, sr + 2, ncols)

    wb.save(path)
    messagebox.showinfo('Thanh cong', f'Da xuat file:\n{path}')
    _open_file(path)


# ── 2. Bang diem lop hoc phan (Giang vien) ────────────────────────────────────

def export_bang_diem_lop(lop_info, rows):
    """
    lop_info : (ma_lop, ten_mon, giang_vien, hoc_ky)
    rows     : list of (id, ma_sv, ho_ten, cc, gk, ck, tb, chu)
    """
    if not OPENPYXL_OK:
        messagebox.showerror('Loi', 'Chua cai openpyxl. Chay: pip install openpyxl'); return

    path = _ask_save(f"BangDiem_Lop_{lop_info[0]}.xlsx")
    if not path: return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Bang diem lop'

    ncols = 9
    ma_lop, ten_mon, gv, hk = lop_info
    _title_row(ws, f'BANG DIEM - {ten_mon.upper()}', ncols, row=1)
    _sub_row(ws, f'Ma lop: {ma_lop}   |   Giang vien: {gv}   |   {hk}', ncols, row=2)

    headers = ['STT', 'Ma SV', 'Ho va ten', 'Chuyen can', 'Giua ky', 'Cuoi ky',
               'Trung binh', 'Xep loai', 'Trang thai']
    _write_header_row(ws, headers, row=4)

    pass_count = 0
    for idx, r in enumerate(rows):
        dr = 5 + idx
        _, ma_sv, ho_ten, cc, gk, ck, tb, chu = r[:8]
        tt = r[8] if len(r) > 8 else '-'
        dat = (tb is not None and tb >= 4.0)
        if dat: pass_count += 1
        vals = [idx+1, ma_sv, ho_ten,
                cc if cc is not None else '-',
                gk if gk is not None else '-',
                ck if ck is not None else '-',
                f'{tb:.2f}' if tb else '-',
                chu or '-', tt or 'Dang hoc']
        _write_data_row(ws, vals, dr, alt=(idx % 2 == 1))
        # Color fail rows
        if tb is not None and tb < 4.0:
            for col in range(1, ncols+1):
                ws.cell(row=dr, column=col).font = Font(name='Calibri', size=10, color='EF4444')

    # Summary
    sr = 5 + len(rows)
    total = len(rows)
    pct = f'{pass_count/total*100:.1f}%' if total else '0%'
    ws.merge_cells(start_row=sr, start_column=1, end_row=sr, end_column=3)
    c = ws.cell(row=sr, column=1,
                value=f'Tong: {total} SV   |   Dat: {pass_count}   |   Rot: {total-pass_count}   |   Ti le dat: {pct}')
    c.font = Font(bold=True, size=10, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='4F46E5')
    c.alignment = _left()
    ws.row_dimensions[sr].height = 22

    # Column widths
    widths = [5, 12, 28, 10, 8, 8, 10, 8, 12]
    for i, w in enumerate(widths, 1):
        _set_col_width(ws, i, w)

    _footer(ws, sr + 2, ncols)
    wb.save(path)
    messagebox.showinfo('Thanh cong', f'Da xuat file:\n{path}')
    _open_file(path)


# ── 3. Danh sach sinh vien (Admin) ────────────────────────────────────────────

def export_danh_sach_sv(rows):
    """rows: list of (id, ma_sv, ho_ten, ngay_sinh, gioi_tinh, ten_lop)"""
    if not OPENPYXL_OK:
        messagebox.showerror('Loi', 'Chua cai openpyxl.'); return

    path = _ask_save('DanhSach_SinhVien.xlsx')
    if not path: return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Danh sach SV'

    ncols = 6
    _title_row(ws, 'DANH SACH SINH VIEN', ncols, row=1)
    _sub_row(ws, f'Tong so: {len(rows)} sinh vien   |   Ngay xuat: {datetime.now().strftime("%d/%m/%Y")}', ncols, row=2)

    headers = ['STT', 'Ma SV', 'Ho va ten', 'Ngay sinh', 'Gioi tinh', 'Lop HC']
    _write_header_row(ws, headers, row=4)

    for idx, r in enumerate(rows):
        dr = 5 + idx
        _, ma_sv, ho_ten, ns, gt, lop = r
        _write_data_row(ws, [idx+1, ma_sv, ho_ten, ns or '-', gt or '-', lop or '-'], dr, alt=(idx%2==1))

    widths = [5, 14, 30, 14, 10, 20]
    for i, w in enumerate(widths, 1):
        _set_col_width(ws, i, w)

    _footer(ws, 6 + len(rows), ncols)
    wb.save(path)
    messagebox.showinfo('Thanh cong', f'Da xuat file:\n{path}')
    _open_file(path)


# ── 4. Danh sach giang vien (Admin) ───────────────────────────────────────────

def export_danh_sach_gv(rows):
    """rows: list of (id, ma_gv, ho_ten, khoa, email, sdt)"""
    if not OPENPYXL_OK:
        messagebox.showerror('Loi', 'Chua cai openpyxl.'); return

    path = _ask_save('DanhSach_GiangVien.xlsx')
    if not path: return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Danh sach GV'

    ncols = 5
    _title_row(ws, 'DANH SACH GIANG VIEN', ncols, row=1)
    _sub_row(ws, f'Tong so: {len(rows)} giang vien   |   Ngay xuat: {datetime.now().strftime("%d/%m/%Y")}', ncols, row=2)

    headers = ['STT', 'Ma GV', 'Ho va ten', 'Khoa', 'Email']
    _write_header_row(ws, headers, row=4)

    for idx, r in enumerate(rows):
        dr = 5 + idx
        _, ma_gv, ho_ten, khoa, email, _ = r
        _write_data_row(ws, [idx+1, ma_gv, ho_ten, khoa or '-', email or '-'], dr, alt=(idx%2==1))

    widths = [5, 14, 30, 24, 28]
    for i, w in enumerate(widths, 1):
        _set_col_width(ws, i, w)

    _footer(ws, 6 + len(rows), ncols)
    wb.save(path)
    messagebox.showinfo('Thanh cong', f'Da xuat file:\n{path}')
    _open_file(path)
